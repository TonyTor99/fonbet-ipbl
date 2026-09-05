"""Выгрузка снимков сборщика киберфутбола FC 26 в Excel.

Строка = матч × 5-минутная отметка игрового времени (плюс строка «до матча»).
Одиночные рынки (1X2, двойные шансы, основная фора, обе забьют) — фикс. колонки.
Тоталы матча и инд. тоталы — ДИНАМИЧЕСКИЕ колонки: для каждой встретившейся линии
свой блок (кф Б / рез Б / кф М / рез М). Результат каждого исхода — рядом с кф.

Запуск:
    python export_cyber.py                 # -> export_cyber_YYYYMMDD_HHMMSS.xlsx
    python export_cyber.py /path/file.xlsx
"""
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import cyber_collector_db as db

# Фикс. колонки: (Заголовок, ключ). Ключи с "__" — спец-обработка ниже.
MAIN_COLUMNS = [
    ("Дата МСК", "__date"),
    ("Время МСК", "__time"),
    ("Лига", "league"),
    ("Команда 1", "team1"),
    ("Команда 2", "team2"),
    ("До матча", "__prematch"),
    ("Отметка мин", "__minute"),
    ("Таймер", "__ts"),
    ("Тайм", "half"),
    ("Счёт тайма", "half_score"),
    ("Таймы", "periods"),
    ("Счёт общий", "__score"),
    # Исход 1X2
    ("П1 кф", "win1_odds"), ("П1", "r_win1"),
    ("X кф", "draw_odds"),  ("X", "r_draw"),
    ("П2 кф", "win2_odds"), ("П2", "r_win2"),
    # Двойные шансы
    ("1X кф", "dc_1x_odds"), ("1X", "r_1x"),
    ("12 кф", "dc_12_odds"), ("12", "r_12"),
    ("X2 кф", "dc_x2_odds"), ("X2", "r_x2"),
    # Фора (основная линия)
    ("Фора К1", "fora_line"), ("Фора К2", "__fora2_line"),
    ("Ф1 кф", "fora1_odds"), ("Ф1", "r_fora1"),
    ("Ф2 кф", "fora2_odds"), ("Ф2", "r_fora2"),
    # Обе забьют
    ("ДА кф", "btts_yes_odds"), ("ДА", "r_btts_yes"),
    ("НЕТ кф", "btts_no_odds"), ("НЕТ", "r_btts_no"),
]

TAIL_COLUMNS = [
    ("Итог счёт", "final_score"),
    ("Итог тотал", "final_total"),
]

KIND_LABEL = {"total": "Тотал", "it1": "ИТ1", "it2": "ИТ2"}

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TOTAL_HEAD_FILL = PatternFill("solid", fgColor="2E7D32")   # зелёная шапка для блоков тоталов
WIN_FILL = PatternFill("solid", fgColor="C6EFCE")
LOSE_FILL = PatternFill("solid", fgColor="FFC7CE")
PUSH_FILL = PatternFill("solid", fgColor="FFEB9C")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RESULT_KEYS = {h for h, k in MAIN_COLUMNS if k and k.startswith("r_")}


def _fmt_line(x: float) -> str:
    return f"{x:g}"


def _value(row: dict, key: str):
    if key == "__score":
        return f"{row['score1']}:{row['score2']}"
    if key == "__prematch":
        return "да" if row.get("is_prematch") else None
    if key == "__minute":
        return "Не начался" if row.get("is_prematch") else row.get("game_minute")
    if key == "__ts":
        ts = row.get("ts")
        return f"{ts // 60}:{ts % 60:02d}" if ts else None
    if key == "__date":
        return (row.get("snap_dt_msk") or "").split(" ")[0] or None
    if key == "__time":
        parts = (row.get("snap_dt_msk") or "").split(" ")
        return parts[1] if len(parts) > 1 else None
    if key == "__fora2_line":
        fl = row.get("fora_line")
        return -fl if fl is not None else None
    return row.get(key)


def _fill_for(val: str):
    if val == "Выигрыш":
        return WIN_FILL
    if val == "Проигрыш":
        return LOSE_FILL
    if val == "Возврат":
        return PUSH_FILL
    return None


def build(path: str):
    rows = db.all_rows()
    lines = db.distinct_lines()   # {kind: [линии]}

    # порядок динамических блоков: total -> it1 -> it2, внутри по возрастанию линии
    total_blocks: list[tuple[str, float]] = []
    for kind in ("total", "it1", "it2"):
        for ln in lines.get(kind, []):
            total_blocks.append((kind, ln))

    # линии тоталов по snapshot_id: {sid: {(kind,line): tl}}
    tl_by_snap: dict[int, dict] = {}
    for tl in db.all_total_lines():
        tl_by_snap.setdefault(tl["snapshot_id"], {})[(tl["kind"], round(tl["line"], 1))] = tl

    wb = Workbook()
    ws = wb.active
    ws.title = "Рынки FC 26"

    # --- шапка ---
    col = 1
    header_meta: list[tuple] = []   # (тип, данные) для каждой колонки
    for head, key in MAIN_COLUMNS:
        header_meta.append(("main", key))
    for kind, ln in total_blocks:
        lbl = f"{KIND_LABEL[kind]} {_fmt_line(ln)}"
        for sub in ("Б кф", "Б", "М кф", "М"):
            header_meta.append(("tl", (kind, round(ln, 1), sub)))
    for head, key in TAIL_COLUMNS:
        header_meta.append(("tail", key))

    # заголовки
    def _head_text(meta) -> str:
        typ, data = meta
        if typ == "tl":
            kind, ln, sub = data
            return f"{KIND_LABEL[kind]} {_fmt_line(ln)} {sub}"
        key = data
        src = MAIN_COLUMNS if typ == "main" else TAIL_COLUMNS
        for h, k in src:
            if k == key:
                return h
        return key

    for c, meta in enumerate(header_meta, 1):
        cell = ws.cell(1, c, _head_text(meta))
        cell.fill = TOTAL_HEAD_FILL if meta[0] == "tl" else HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = "A2"

    # --- строки ---
    for r, row in enumerate(rows, 2):
        snap_tls = tl_by_snap.get(row["id"], {})
        for c, meta in enumerate(header_meta, 1):
            typ, data = meta
            if typ == "main":
                val = _value(row, data)
                cell = ws.cell(r, c, val)
                head = _head_text(meta)
                if head in RESULT_KEYS and val:
                    fill = _fill_for(val)
                    if fill:
                        cell.fill = fill
                    cell.alignment = Alignment(horizontal="center")
            elif typ == "tail":
                cell = ws.cell(r, c, _value(row, data))
            else:  # tl
                kind, ln, sub = data
                tl = snap_tls.get((kind, ln))
                val = None
                is_res = False
                if tl:
                    if sub == "Б кф":
                        val = tl.get("b_odds")
                    elif sub == "М кф":
                        val = tl.get("m_odds")
                    elif sub == "Б":
                        val = tl.get("r_b"); is_res = True
                    elif sub == "М":
                        val = tl.get("r_m"); is_res = True
                cell = ws.cell(r, c, val)
                if is_res and val:
                    fill = _fill_for(val)
                    if fill:
                        cell.fill = fill
                    cell.alignment = Alignment(horizontal="center")
            cell.border = BORDER

    # --- ширины ---
    for c, meta in enumerate(header_meta, 1):
        letter = ws.cell(1, c).column_letter
        head = _head_text(meta)
        ws.column_dimensions[letter].width = max(8, min(20, len(head) + 2))

    wb.save(path)
    st = db.stats()
    print(f"Сохранено: {path}")
    print(f"Строк: {st['rows']} | матчей: {st['events']} | с результатом: {st['resolved']}")
    print(f"Линий тоталов в колонках: {len(total_blocks)} ({total_blocks})")


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else \
        f"export_cyber_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    db.init_db()
    build(path)


if __name__ == "__main__":
    main()
