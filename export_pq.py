"""Выгрузка ОТПРАВЛЕННЫХ сигналов стратегии ЧЕТВЕРТИ Pro Жен в Excel
(из pq_women_strategy.db). Строка = отправленный сигнал (ТБ/ТМ четверти).

Запуск:
    python export_pq.py                 # всё
    python export_pq.py file.xlsx
"""
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import pq_db
from config import PQ_SIDES

COLUMNS = [
    ("Дата МСК", "__date"),
    ("Время МСК", "__time"),
    ("Команда 1", "team1"),
    ("Команда 2", "team2"),
    ("Сторона", "__side"),
    ("Минута", "minute"),
    ("Четверть", "quarter"),
    ("Линия", "line"),
    ("Кф", "odds"),
    ("Счёт матча (сигнал)", "__score"),
    ("Счёт четв. (сигнал)", "q_live"),
    ("Итог четв.", "q_final"),
    ("Тотал четв.", "q_total"),
    ("Результат", "result"),
    ("Прибыль", "profit"),
]

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WIN_FILL = PatternFill("solid", fgColor="C6EFCE")
LOSE_FILL = PatternFill("solid", fgColor="FFC7CE")
PUSH_FILL = PatternFill("solid", fgColor="FFEB9C")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _value(row: dict, key: str):
    if key == "__side":
        return PQ_SIDES.get(row.get("side"), row.get("side"))
    if key == "__score":
        return f"{row['score1']}:{row['score2']}"
    if key == "__date":
        return (row.get("created_at") or "").split(" ")[0] or None
    if key == "__time":
        parts = (row.get("created_at") or "").split(" ")
        return parts[1] if len(parts) > 1 else None
    return row.get(key)


def build(path: str, title: str = "Сигналы Четверти ProЖ") -> int:
    rows = pq_db.signals_for_export()
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    for c, (head, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(1, c, head)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = "A2"

    for r, row in enumerate(rows, 2):
        result = row.get("result")
        for c, (head, key) in enumerate(COLUMNS, 1):
            val = _value(row, key)
            cell = ws.cell(r, c, val)
            cell.border = BORDER
            if head == "Результат" and val:
                if result == "Выигрыш":
                    cell.fill = WIN_FILL
                elif result == "Проигрыш":
                    cell.fill = LOSE_FILL
                elif result == "Возврат":
                    cell.fill = PUSH_FILL
                cell.alignment = Alignment(horizontal="center")

    for c, (head, _) in enumerate(COLUMNS, 1):
        letter = ws.cell(1, c).column_letter
        ws.column_dimensions[letter].width = max(9, min(22, len(head) + 2))

    wb.save(path)
    return len(rows)


def main():
    pq_db.init_db()
    path = sys.argv[1] if len(sys.argv) > 1 else \
        f"export_pq_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    n = build(path)
    print(f"Сохранено: {path} | строк: {n}")


if __name__ == "__main__":
    main()
